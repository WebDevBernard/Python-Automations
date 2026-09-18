import math
import sys
from datetime import datetime
from pathlib import Path
import openpyxl
from utils import progressbar, write_to_new_docx
from constants import get_insurer, resolve_insurer

DATE_FORMAT = "%B %d, %Y"

BASE_DIR = Path(__file__).resolve().parent
CONFIG_PATH = BASE_DIR / "config.xlsx"

try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass


def safe_strip(value):
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return str(value).strip()


def smart_title(text):
    """
    Title-cases words, but leaves a word untouched if:
      - it contains a digit, or
      - the NEXT word starts with a digit (so labels like "BCS 3746" stay intact)
    Certain small words (e.g. "of") are kept lowercase unless they're the first word.
    """
    if not text:
        return text

    lowercase_words = {"of", "the", "and", "a", "an", "in", "on", "for"}

    words = text.split(" ")
    result = []
    for i, word in enumerate(words):
        has_digit = any(ch.isdigit() for ch in word)
        next_word = words[i + 1] if i + 1 < len(words) else ""
        next_starts_digit = bool(next_word) and next_word[0].isdigit()

        if has_digit or next_starts_digit:
            result.append(word)
        elif word.lower() in lowercase_words and i != 0:
            result.append(word.lower())
        else:
            result.append(word[:1].upper() + word[1:].lower() if word else word)
    return " ".join(result)


def to_float(value):
    if value is None:
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace("$", "").replace(",", "")
    if not s:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0


def format_amount(value):
    """Strips any existing $ / commas and reformats as a plain number string."""
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return f"{value:,.2f}"
    s = str(value).strip().replace("$", "").replace(",", "")
    if not s:
        return ""
    try:
        return f"{float(s):,.2f}"
    except ValueError:
        return s


def parse_date(value):
    if not value:
        return ""
    if isinstance(value, datetime):
        return value.strftime(DATE_FORMAT)
    value = str(value).strip()
    date_formats = [
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%B %d, %Y",
        "%b %d, %Y",
        "%d-%b-%y",
    ]
    for fmt in date_formats:
        try:
            return datetime.strptime(value, fmt).strftime(DATE_FORMAT)
        except ValueError:
            continue
    print(f"\u26a0\ufe0f Could not parse date: {value}")
    return value


def parse_date_to_dt(value):
    """Same parsing as parse_date, but returns a datetime object (or None)."""
    if not value:
        return None
    if isinstance(value, datetime):
        return value
    value = str(value).strip()
    date_formats = [
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%m/%d/%Y",
        "%B %d, %Y",
        "%b %d, %Y",
        "%d-%b-%y",
    ]
    for fmt in date_formats:
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    print(f"\u26a0\ufe0f Could not parse transaction date: {value}")
    return None


def load_producer_mapping(mapping_path):
    wb = openpyxl.load_workbook(mapping_path, data_only=True)
    if "File Completion Tool" not in wb.sheetnames:
        return {}
    ws = wb["File Completion Tool"]
    mapping = {}
    row = 27
    while True:
        code = safe_strip(ws.cell(row=row, column=1).value)
        name = safe_strip(ws.cell(row=row, column=2).value)
        if not code and not name:
            break
        if code:
            mapping[code.lower()] = name
        row += 1
    return mapping


def load_transaction_years(mapping_path):
    """
    Reads the 'Transactions' worksheet and builds a mapping of:
        policy_number (normalized) -> { year: {"date": datetime, "amount": float,
                                                "risk_location": str, "old_gore_policy": str} }

    Only rows where transaction type is 'New', 'Renewal', or 'Re-write' are included.
    Rows sharing the same policy AND the same exact date have their
    amounts summed together (treated as one transaction).
    If a policy has multiple distinct dates within the same year,
    the latest date's group wins for that year.
    If a 'risk_location' column exists and has a value, it overrides the
    mailing address for that transaction.
    If an 'old_gore_policy' column exists and has a value, it replaces the
    policy number used for the disclosure letter output.
    """
    wb = openpyxl.load_workbook(mapping_path, data_only=True)
    if "Transactions" not in wb.sheetnames:
        print("\u274c 'Transactions' sheet not found in config.xlsx")
        return {}

    ws = wb["Transactions"]

    headers = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        header = safe_strip(cell.value).lower()
        if header:
            headers[header] = col_idx

    required = ["date", "transaction", "policy", "amount"]
    missing = [c for c in required if c not in headers]
    if missing:
        print(f"\u274c Missing required columns in Transactions sheet: {missing}")
        return {}

    has_risk_location = "risk_location" in headers
    has_old_gore_policy = "old_gore_policy" in headers
    valid_types = {"new", "renewal", "rewrite"}

    # Step 1: group by (policy, exact date) and sum amounts
    date_groups = (
        {}
    )  # (policy_key, date_obj) -> {"date": dt, "amount": float, "risk_location": str}

    for row in ws.iter_rows(min_row=2, values_only=False):
        policy_raw = safe_strip(row[headers["policy"] - 1].value)
        txn_type_raw = safe_strip(row[headers["transaction"] - 1].value).lower()
        txn_type = txn_type_raw.replace("-", "").replace(" ", "")
        date_val = row[headers["date"] - 1].value
        amount_val = row[headers["amount"] - 1].value

        if not policy_raw or txn_type not in valid_types:
            continue

        dt = parse_date_to_dt(date_val)
        if dt is None:
            continue

        risk_location_val = ""
        if has_risk_location:
            risk_location_val = safe_strip(row[headers["risk_location"] - 1].value)
        old_gore_policy_val = ""
        if has_old_gore_policy:
            old_gore_policy_val = safe_strip(row[headers["old_gore_policy"] - 1].value)

        policy_key = policy_raw.strip().upper()
        group_key = (policy_key, dt.date())

        if group_key not in date_groups:
            date_groups[group_key] = {
                "date": dt,
                "amount": 0.0,
                "risk_location": "",
                "old_gore_policy": "",
            }

        date_groups[group_key]["amount"] += to_float(amount_val)

        # Keep the risk_location if we find a non-empty one (last non-empty wins)
        if risk_location_val:
            date_groups[group_key]["risk_location"] = risk_location_val
        # Keep the old_gore_policy if we find a non-empty one (last non-empty wins)
        if old_gore_policy_val:
            date_groups[group_key]["old_gore_policy"] = old_gore_policy_val

    # Step 2: roll grouped entries up by year (latest date in the year wins if there
    # happen to be multiple distinct dates for the same policy in the same year)
    policy_years = {}
    for (policy_key, _date_obj), group in sorted(
        date_groups.items(), key=lambda kv: kv[1]["date"]
    ):
        year = group["date"].year
        policy_years.setdefault(policy_key, {})
        policy_years[policy_key][year] = group

    return policy_years


def read_sheet_rows(mapping_path, producer_mapping):
    wb = openpyxl.load_workbook(mapping_path, data_only=True)
    if "Disclosure Notice" not in wb.sheetnames:
        print("\u274c 'Disclosure Notice' sheet not found in config.xlsx")
        return []

    ws = wb["Disclosure Notice"]

    headers = {}
    for col_idx, cell in enumerate(ws[1], start=1):
        header = safe_strip(cell.value).lower()
        if header:
            headers[header] = col_idx

    required = [
        "policynum",
        "name",
        "h_address1",
        "h_cityprov",
        "h_postzip",
        "effective",
    ]
    missing = [c for c in required if c not in headers]
    if missing:
        print(f"\u274c Missing required columns in Disclosure Notice sheet: {missing}")
        return []

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=False):
        if (
            row[headers["policynum"] - 1].value is None
            and row[headers["name"] - 1].value is None
        ):
            continue
        policynum = safe_strip(row[headers["policynum"] - 1].value)
        name = safe_strip(row[headers["name"] - 1].value)
        if not policynum and not name:
            continue

        entry = {
            "policy_number": policynum,
            "named_insured": smart_title(name),
            "address_line_one": safe_strip(row[headers["h_address1"] - 1].value),
            "address_line_two": safe_strip(row[headers["h_cityprov"] - 1].value),
            "address_line_three": safe_strip(row[headers["h_postzip"] - 1].value),
            "effective_date": row[headers["effective"] - 1].value,
        }

        for col_name in ("insurer", "prem_amt"):
            if col_name in headers:
                entry[col_name] = safe_strip(row[headers[col_name] - 1].value)
            else:
                entry[col_name] = ""

        pcode_raw = ""
        if "pcode" in headers:
            pcode_raw = safe_strip(row[headers["pcode"] - 1].value)
        entry["producer_name"] = producer_mapping.get(pcode_raw.lower(), "")

        rows.append(entry)

    return rows


def disclosure_notice(config=None):
    mapping_path = CONFIG_PATH

    producer_mapping = load_producer_mapping(mapping_path)
    rows = read_sheet_rows(mapping_path, producer_mapping)
    if not rows:
        print("No data found in Disclosure Notice sheet")
        return

    transaction_years = load_transaction_years(mapping_path)
    if not transaction_years:
        print("No New/Renewal/Re-write transactions found in Transactions sheet")
        return

    base_output_dir = Path.home() / "Desktop" / "Strata Disclosure Notices"
    base_output_dir.mkdir(parents=True, exist_ok=True)

    success_count = 0
    total_expected = 0

    for row_data in progressbar(rows, prefix="Generating letters "):
        policy_key = row_data["policy_number"].strip().upper()
        years_for_policy = transaction_years.get(policy_key)

        if not years_for_policy:
            print(
                f"\u26a0\ufe0f No matching New/Renewal/Re-write transaction found for policy "
                f"{row_data['policy_number']} ({row_data.get('named_insured', 'unknown')}) \u2014 skipped"
            )
            continue

        name_dir = base_output_dir / row_data["named_insured"]
        name_dir.mkdir(parents=True, exist_ok=True)

        row_data["insurer"] = resolve_insurer(
            row_data.get("insurer", ""), row_data["policy_number"]
        )

        for year, txn_info in sorted(years_for_policy.items()):
            try:
                row_data["effective_date"] = parse_date(txn_info["date"])
                row_data["prem_amt"] = format_amount(txn_info["amount"])

                if txn_info.get("risk_location"):
                    row_data["mailing_address"] = txn_info["risk_location"]
                else:
                    row_data["mailing_address"] = "\n".join(
                        part
                        for part in (
                            row_data.get("address_line_one", ""),
                            row_data.get("address_line_two", ""),
                            row_data.get("address_line_three", ""),
                        )
                        if part
                    )

                if txn_info.get("old_gore_policy"):
                    row_data["policy_number"] = txn_info["old_gore_policy"]

                template = BASE_DIR / "assets" / "Strata Disclosure Notice.docx"
                if write_to_new_docx(
                    template_path=template,
                    data=row_data,
                    output_dir=name_dir,
                    output_suffix="Disclosure Notice",
                    output_year=f"({year})",
                ):
                    success_count += 1
                total_expected += 1
            except Exception as e:
                import traceback

                print(
                    f"\u274c Failed for {row_data.get('named_insured', 'unknown')} ({year}): {e}"
                )
                traceback.print_exc()

    print(
        f"******** Disclosure Notice completed: {success_count}/{total_expected} letters generated ********"
    )


if __name__ == "__main__":
    disclosure_notice()
